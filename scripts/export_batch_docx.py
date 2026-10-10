# -*- coding: utf-8 -*-
"""
Script tự động kết xuất trọn bộ 5 File Word (.docx) chuyên dụng cho sản xuất hàng loạt
Skill: han-daily-family-vlog
Quy chuẩn: 100% văn bản thuần túy (Plain Text Paragraphs), KHÔNG sử dụng Table ở bất kỳ file nào.
Mục tiêu: Đảm bảo copy-paste 1 chạm siêu tốc vào các AI tools (Midjourney, Flux, Kling, Luma, CapCut).
"""

import os
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_styled_title(doc, title_text, subtitle_text, note_text=None):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run(title_text)
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(234, 88, 12) # Warm playful orange
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    if note_text:
        p_sub.paragraph_format.space_after = Pt(4)
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_note = p_note.add_run(note_text)
        run_note.font.name = "Segoe UI"
        run_note.font.size = Pt(9.5)
        run_note.font.color.rgb = RGBColor(225, 29, 72)
        p_note.paragraph_format.space_after = Pt(16)
    else:
        p_sub.paragraph_format.space_after = Pt(14)

def add_prompt_paragraph(doc, text_content, label=None, font_name="Consolas", font_size=10, is_bold=False):
    if label:
        p_label = doc.add_paragraph()
        run_lbl = p_label.add_run(label)
        run_lbl.font.name = "Segoe UI"
        run_lbl.font.size = Pt(10.5)
        run_lbl.font.bold = True
        run_lbl.font.color.rgb = RGBColor(234, 88, 12)
        p_label.paragraph_format.space_before = Pt(8)
        p_label.paragraph_format.space_after = Pt(2)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text_content)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = is_bold
    run.font.color.rgb = RGBColor(30, 41, 59)

def export_all_docx(project_dir, project_data):
    """
    Kết xuất 5 file Word theo chuẩn 100% Plain Text Paragraphs không Table
    """
    export_dir = os.path.join(project_dir, "export_batch")
    os.makedirs(export_dir, exist_ok=True)

    # 1. Kịch bản tổng thể
    doc1 = Document()
    add_styled_title(doc1, f"KỊCH BẢN TỔNG THỂ: {project_data.get('title')}", project_data.get('subtitle', 'Hân Daily Family Vlog'))
    for sc in project_data.get('scenes', []):
        add_prompt_paragraph(doc1, 
            f"Thời lượng: {sc.get('duration')}\n"
            f"Bối cảnh & Tình huống diễn xuất: {sc.get('description', '')}\n"
            f"Góc máy & Chuyển động: {sc.get('camera', '')}\n"
            f"Lời thoại / Lồng tiếng: {sc.get('dialogue', '')}",
            label=f"PHÂN CẢNH {sc.get('id')}: {sc.get('name')}",
            font_name="Segoe UI", font_size=10.5
        )
    doc1.save(os.path.join(export_dir, "01_Kich_Ban_Tong_The.docx"))

    # 2. Batch Prompt Image (Flux / Midjourney)
    doc2 = Document()
    add_styled_title(doc2, "PROMPT TẠO ẢNH HÀNG LOẠT (FLUX / MIDJOURNEY)", "Bộ Prompt Keyframe 9:16 chuẩn Real Human DSLR, Cute Toddler Expressions & Family Turnaround", "💡 HƯỚNG DẪN: Copy từng đoạn prompt bên dưới dán trực tiếp vào Midjourney hoặc Flux. Dùng kèm ảnh Character Turnaround tương ứng.")
    for sc in project_data.get('scenes', []):
        add_prompt_paragraph(doc2, sc.get('image_prompt', ''), label=f"CẢNH {sc.get('id')} - {sc.get('name')}")
    doc2.save(os.path.join(export_dir, "02_Prompt_Tao_Anh_Image_Batch.docx"))

    # 3. Batch Prompt Video (Kling / Luma / Hailuo / Wan)
    doc3 = Document()
    add_styled_title(doc3, "PROMPT TẠO VIDEO CHUYỂN ĐỘNG & DIỄN XUẤT (KLING / LUMA / HAILUO / WAN)", "Bộ Prompt Chuyển Động Biểu Cảm Má Bánh Bao, Cử Chỉ Ngây Ngô & Diễn Xuất Gia Đình", "💡 HƯỚNG DẪN: Upload Keyframe ảnh từ File 02 làm Start Frame (I2V), copy prompt video và thiết lập motion tương ứng.")
    for sc in project_data.get('scenes', []):
        content = f"{sc.get('video_prompt', '')}\n\n[Motion Settings]: {sc.get('motion_params', '')}"
        add_prompt_paragraph(doc3, content, label=f"CẢNH {sc.get('id')} - {sc.get('name')}")
    doc3.save(os.path.join(export_dir, "03_Prompt_Tao_Video_Motion_Batch.docx"))

    # 4. Âm thanh, Thoại Em Bé & SFX Hài Hước
    doc4 = Document()
    add_styled_title(doc4, "KỊCH BẢN THOẠI EM BÉ, FOLEY & HIỆU ỨNG ÂM THANH HÀI HƯỚNG", "Kịch bản lồng tiếng em bé ngọng nghịu, tiếng cười nắc nẻ, tiếng dỗi hờn và nhạc nền Acoustic/Piano", "💡 HƯỚNG DẪN: Dùng CapCut hoặc Premiere ghép audio voiceover khớp đúng khẩu hình diễn xuất của bé và ba mẹ.")
    for sc in project_data.get('scenes', []):
        content = f"Lời thoại nhân vật: {sc.get('dialogue', '')}\nHiệu ứng âm thanh hài (SFX): {sc.get('sfx', '')}\nNhạc nền BGM: {sc.get('bgm', '')}"
        add_prompt_paragraph(doc4, content, label=f"CẢNH {sc.get('id')} - {sc.get('name')}", font_name="Segoe UI", font_size=10.5)
    doc4.save(os.path.join(export_dir, "04_Am_Thanh_Voice_BGM.docx"))

    # 5. Lệnh FFmpeg Ghép Video Tự Động
    doc5 = Document()
    add_styled_title(doc5, "LỆNH FFMPEG GHÉP VIDEO TỰ ĐỘNG", "Command line render chuẩn tỷ lệ 9:16 (1080x1920), 60fps mượt mà", "💡 HƯỚNG DẪN: Mở terminal tại thư mục chứa các clip MP4 đã render từ AI, dán lệnh để xuất bản video hoàn chỉnh.")
    for idx, cmd in enumerate(project_data.get('ffmpeg_commands', []), 1):
        add_prompt_paragraph(doc5, cmd, label=f"Lệnh {idx}")
    doc5.save(os.path.join(export_dir, "05_Lenh_FFmpeg_Xuat_Video.docx"))

    print(f"[OK] Da ket xuat thanh cong tron bo 5 file .docx khong co Table tai: {export_dir}")
    return export_dir

if __name__ == "__main__":
    print("Script export_batch_docx.py cho skill han-daily-family-vlog da san sang.")
