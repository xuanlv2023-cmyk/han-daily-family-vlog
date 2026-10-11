#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Script tự động kết xuất trọn bộ 5 File Word (.docx) chuyên dụng cho sản xuất hàng loạt
Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại (Clone Linh Mít Video 1)
Skill: han-daily-family-vlog
Khung phân cảnh: Chuẩn 2 phân cảnh x 10.0s = 20.0s cho Google Veo 3.1 & Muse
1. 00_HO_SO_THAM_CHIEU_GOC.docx
2. 01_PROMPTS_ANH_HANG_LOAT.docx
3. 02_PROMPTS_VIDEO_HANG_LOAT.docx
4. 03_PROMPTS_VOICE_HANG_LOAT.docx
5. 04_METADATA_SEO.docx
"""

import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_styled_title(doc, title_text, subtitle_text, note_text=None):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run(title_text)
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(234, 88, 12) # Warm Orange Family
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = p_sub.add_run(subtitle_text)
    run_sub.font.name = 'Segoe UI'
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    if note_text:
        p_sub.paragraph_format.space_after = Pt(4)
        p_note = doc.add_paragraph()
        p_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run_note = p_note.add_run(note_text)
        run_note.font.name = 'Segoe UI'
        run_note.font.size = Pt(9.5)
        run_note.font.color.rgb = RGBColor(194, 65, 12)
        p_note.paragraph_format.space_after = Pt(16)
    else:
        p_sub.paragraph_format.space_after = Pt(14)

def add_clean_paragraph(doc, text_content, font_name='Consolas', font_size=10, is_bold=False):
    """Thêm đúng 1 đoạn văn liên tục không nhãn thừa cho tool batch"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text_content)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = is_bold
    run.font.color.rgb = RGBColor(30, 41, 59)

def export_file_00(output_dir):
    doc = Document()
    add_styled_title(doc, "HỒ SƠ THAM CHIẾU GỐC (MASTER REFERENCE ASSETS)",
                     "Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại • 3 Nhân Vật & Nhà Hàng",
                     "⚠️ BƯỚC 0 BẮT BUỘC: Nạp ảnh Turnaround làm Character Reference cho Tool AI tạo ảnh/video.")

    p_h1 = doc.add_heading(level=2)
    p_h1.add_run("1. Master Toddler Model Sheet @_Brand_Baby_Lead (Bé Hân)")
    prompt_baby = "A professional four-view character turnaround model sheet of @_Brand_Baby_Lead, an authentic 3-year-old Vietnamese toddler girl, displaying front view, three-quarter view, profile side view, and close-up dramatic sulking pout expression with puffed dumpling cheeks and sideways side-eye glare. Super adorable chubby cheeks, expressive dark eyes, topknot bun with pearl daisy hairpin, natural toddler skin with delicate pores. Wearing cream knit dress and navy blue cardigan with white collar, 5-finger hands. Soft daylight, neutral studio background, raw DSLR photography --ar 16:9 --style raw"
    add_clean_paragraph(doc, prompt_baby)

    p_h2 = doc.add_heading(level=2)
    p_h2.add_run("2. Master Father Model Sheet @_Brand_Father_Lead (Ba Hoàng Long)")
    prompt_dad = "A professional turnaround model sheet of @_Brand_Father_Lead, a 27-year-old Vietnamese young father, round black glasses, olive sweater, warm playful smile, real human skin pores, Sony A7 IV photography --ar 16:9 --style raw"
    add_clean_paragraph(doc, prompt_dad)

    p_h3 = doc.add_heading(level=2)
    p_h3.add_run("3. Master Mother Model Sheet @_Brand_Mother_Lead (Mẹ Uyên Hân)")
    prompt_mom = "A professional turnaround model sheet of @_Brand_Mother_Lead, a 25-year-old Vietnamese stylish mother, dark wavy hair, navy knit top, natural gentle beauty, raw photography --ar 16:9 --style raw"
    add_clean_paragraph(doc, prompt_mom)

    p_h4 = doc.add_heading(level=2)
    p_h4.add_run("4. Master Location Sheet @_Location_Restaurant_Dining")
    prompt_loc = "A cinematic establishing photo of @_Location_Restaurant_Dining, a warm modern family restaurant dining room, white marble dining table with plate of crispy golden fried chicken nuggets, warm diffused ambient ceiling spotlights, natural indoor shadows, 9:16 vertical ratio, clean aesthetic sensor image, zero watermark --ar 9:16 --style raw"
    add_clean_paragraph(doc, prompt_loc)

    doc.save(os.path.join(output_dir, "00_HO_SO_THAM_CHIEU_GOC.docx"))

def export_file_01(output_dir):
    doc = Document()
    add_styled_title(doc, "PROMPTS TẠO ẢNH HÀNG LOẠT (BATCH IMAGE PROMPTS)",
                     "Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại • Chuẩn 9:16 Real Human Photography",
                     "Định dạng Tool Batch: 100% văn bản thuần túy không nhãn thừa. Mỗi cảnh là đúng 1 đoạn văn liên tục xuống dòng (2 cảnh = đúng 2 đoạn văn sạch).")

    prompts = [
        "Authentic raw 8k DSLR candid photograph of a super cute 3-year-old Vietnamese girl (Bé Hân) with chubby rosy dumpling cheeks and high bun with pearl daisy hairpin, sitting in the center of a dining table between her stylish mother in navy knit top and handsome father wearing round black glasses in olive shirt. On the white marble table is a plate of crispy golden popcorn chicken. The toddler has an iconic hilarious dramatic sulking face with puffed dumpling cheeks, exaggerated lower lip pouting downward, and eyes giving an epic sideways side-eye glare at her smiling father, natural child skin pores, warm restaurant indoor lighting, 9:16 vertical ratio, shot on Sony A7R V with 85mm f/1.4 lens, ultra photorealistic, zero CGI, zero plastic skin --ar 9:16 --v 6.1 --style raw",
        "Authentic raw 8k DSLR candid photograph of Bé Hân turning toward her laughing mother with dramatic teary pouting face and crossed tiny arms, complaining about her stolen chicken nugget, while father chuckles playfully in background, cozy restaurant atmosphere, crystal-sharp baby facial texture, realistic child tears on rosy cheeks, Sony A7R V, 85mm f/1.4, raw photography --ar 9:16 --v 6.1 --style raw"
    ]
    for p in prompts:
        add_clean_paragraph(doc, p)
    doc.save(os.path.join(output_dir, "01_PROMPTS_ANH_HANG_LOAT.docx"))

def export_file_02(output_dir):
    doc = Document()
    add_styled_title(doc, "PROMPTS TẠO VIDEO HÀNG LOẠT (BATCH MEGA PROMPTS)",
                     "Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại • Chuẩn Khung 10.0s Cho Google Veo 3.1 & Muse",
                     "Định dạng Tool Batch: 100% văn bản thuần túy không nhãn thừa. Mỗi cảnh là đúng 1 đoạn văn liên tục không ngắt dòng bên trong (2 cảnh = đúng 2 đoạn văn sạch).")

    video_prompts = [
        "[SPECIFICATION]: Exactly 10.0 seconds, 9:16 vertical 1080x1920, 30 FPS, authentic candid family snapshot, unretouched real human photography, zero AI plastic look, zero CGI rendering. [CAMERA LOCK & FRAMING]: Locked-off eye-level medium shot framing all three family members around the dining table, stationary tripod-mounted camera. Continuous unbroken single-take video throughout all 10.0 seconds with zero camera movement, no zoom, no pan, strictly no scene cut. [FAMILY CHARACTER ANCHORS]: Center: @_Brand_Baby_Lead (super cute 3-year-old Vietnamese girl, chubby rosy dumpling cheeks, expressive dark brown eyes, high topknot bun with pearl daisy hairpin, authentic soft baby skin with natural micro-pores). Left: @_Brand_Mother_Lead (25-year-old stylish mother with long wavy dark hair, elegant fair complexion). Right: @_Brand_Father_Lead (27-year-old handsome father with round black-rimmed glasses, playful smile). [DYNAMIC OUTFITS]: @_Brand_Baby_Lead wearing cream knitted dress and navy blue cardigan with white collar, mother wearing navy knit polo top with white collar, father wearing olive-green long-sleeve tee, clean anatomically perfect 5-finger hands. [ENVIRONMENT & LIGHTING]: Positioned in @_Location_Restaurant_Dining, white marble dining table with plate of crispy golden popcorn chicken in foreground, cozy restaurant background with warm ambient ceiling spotlights, natural soft contact shadows. [INITIAL STATE]: Toddler sitting in highchair between parents, mouth wide open awaiting food, blowing gently with tiny puckered lips, eyes sparkling with innocent anticipation. [COMEDY SKIT & REACTION TIMELINE]: - [0.0s - 3.0s]: Father holds stainless steel fork with a piece of crispy chicken extending toward toddler's mouth, toddler opens mouth wide blowing innocently with puffed cheeks. - [3.0s - 7.0s]: Sudden comedic twist, father smoothly pivots the fork left directly into mother's mouth, mother bites the chicken happily and chews with joyful satisfaction; toddler freezes completely with mouth still wide open in disbelief. - [7.0s - 10.0s]: Dumpling cheeks slowly puff up like steamed buns, exaggerated lower lip pouts downward in immense grievance, eyes shift sideways into an epic hilarious side-eye glare at father without moving her head, settling into Keyframe K2. [AUDIO DIALOGUE & BABY LIP-SYNC]: Father (0-3s): \"Nguội rồi nè, há miệng to ba đút cho công chúa nhen!\", Mother (4-6s): \"Ưm giòn ngon quá ba ơi!\", cartoon crow caw SFX during freeze, funny spring boing SFX as cheeks puff out. [FINAL STATE]: Toddler maintaining iconic frozen side-eye glare, parents holding satisfied chuckles, matching Keyframe K2. [NEGATIVE CONSTRAINTS]: Strictly NO extra fingers, NO fused hands, NO distorted baby limbs, NO waxy plastic doll face, NO 3D Pixar anime look, NO text overlays, NO watermarks.",
        "[SPECIFICATION]: Exactly 10.0 seconds, 9:16 vertical 1080x1920, 30 FPS, candid emotional toddler comedy, unretouched real human photography, zero AI plastic look. [CAMERA LOCK & FRAMING]: Locked-off medium close-up shot focused on Bé Hân, father and mother visible on both sides smiling warmly, perfectly stabilized stationary camera. Continuous unbroken single take. [FAMILY CHARACTER ANCHORS]: @_Brand_Baby_Lead in sharp focus with puffed rosy dumpling cheeks and misty expressive eyes; @_Brand_Mother_Lead and @_Brand_Father_Lead interacting lovingly. [DYNAMIC OUTFITS]: Identical outfits: cream knit dress with blue cardigan, navy knit top, olive long-sleeve tee, clean 5 fingers. [ENVIRONMENT & LIGHTING]: Restaurant dining setting @_Location_Restaurant_Dining, warm indoor lighting, plate of golden chicken on marble table. [INITIAL STATE]: Seamlessly inherits Keyframe K2: toddler in dramatic sideways glare, puffed dumpling cheeks. [COMEDY SKIT & REACTION TIMELINE]: - [0.0s - 3.0s]: Toddler slowly breaks her frozen glare, turning her head toward mother with exaggerated pouting lips and teary indignant eyes, pointing tiny finger at empty fork. - [3.0s - 7.0s]: Delivers her hilarious baby scolding line with pouting baby lip-sync: \"Mẹ kì zậy... con thổi nguội gòi mà mẹ ăn hết trơn!\", crossing tiny arms in theatrical grievance. - [7.0s - 10.0s]: Father and mother burst into uncontrollable genuine contagious laughter, father reaching out to gently pat toddler's head while mother hugs her lovingly to comfort her. [AUDIO DIALOGUE & BABY LIP-SYNC]: Toddler (3-7s ngọng líu dỗi hờn): \"Mẹ kì zậy... con thổi nguội gòi mà mẹ ăn hết trơn!\", Father and Mother (7-10s): (bursting into authentic warm contagious laughter), cheerful acoustic ukulele bed at -22dB. [FINAL STATE]: Toddler nestled in mother's arms while parents laugh warmly, completing the video loop. [NEGATIVE CONSTRAINTS]: Strictly NO extra fingers, NO morphing eyes, NO AI plastic skin, NO 3D cartoon style, NO watermarks."
    ]
    for vp in video_prompts:
        add_clean_paragraph(doc, vp)
    doc.save(os.path.join(output_dir, "02_PROMPTS_VIDEO_HANG_LOAT.docx"))

def export_file_03(output_dir):
    doc = Document()
    add_styled_title(doc, "KỊCH BẢN LỜI THOẠI HÀNG LOẠT (BATCH TTS SCRIPTS)",
                     "Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại • Chuẩn 2 Cảnh x 10.0s • 100% Giọng Miền Nam",
                     "Định dạng Tool Batch: Lời thoại tiếng Việt sạch 100%, không nhãn người nói, không chú thích SFX. Mỗi cảnh là đúng 1 đoạn văn liên tục xuống dòng (2 cảnh = đúng 2 đoạn văn sạch).")

    dialogue_lines = [
        "Nguội rồi nè, há miệng to ba đút cho công chúa nhen! Ưm giòn ngon quá ba ơi!",
        "Mẹ kì zậy... con thổi nguội gòi mà mẹ ăn hết trơn!"
    ]
    for dl in dialogue_lines:
        add_clean_paragraph(doc, dl, font_name='Segoe UI', font_size=11.5)
    doc.save(os.path.join(output_dir, "03_PROMPTS_VOICE_HANG_LOAT.docx"))

def export_file_04(output_dir):
    doc = Document()
    add_styled_title(doc, "GÓI TỐI ƯU SEO & ĐĂNG TẢI VIRAL ĐA KÊNH",
                     "Dự án: Cú Lừa Gà Rán & Pha Liếc Xéo Huyền Thoại • Chuẩn Đăng Tải Shorts / TikTok / Reels")

    p_h1 = doc.add_heading(level=2)
    p_h1.add_run("1. Prompt Tạo Ảnh Bìa Thumbnail (9:16)")
    thumb_prompt = "A hilarious viral family thumbnail photo, 9:16 vertical, close-up on @_Brand_Baby_Lead (Bé Hân 3 tuổi) with iconic dramatic sulking pout face, cheeks puffed like steamed dumplings, giving an epic side-eye glare at her smiling parents at the restaurant table, golden fried chicken plate in foreground, bright warm lighting, crisp DSLR depth of field, 8k raw photography, zero plastic AI smoothing --ar 9:16 --v 6.1 --style raw"
    add_clean_paragraph(doc, thumb_prompt)

    p_h2 = doc.add_heading(level=2)
    p_h2.add_run("2. Chữ Dán Thumbnail (Overlay Text)")
    add_clean_paragraph(doc, "CÚ LỪA GÀ RÁN! 🍗😱\nPHA LIẾC HUYỀN THOẠI! 👀🤣\nDỖI LUÔN KHỎI ĂN! 😤💔", font_name='Segoe UI', font_size=11)

    p_h3 = doc.add_heading(level=2)
    p_h3.add_run("3. Tiêu Đề Đa Nền Tảng (A/B Testing Headlines)")
    add_clean_paragraph(doc, "A: Ba lừa đút gà rán cho mẹ và pha liếc xéo huyền thoại của bé Hân 🍗👀 #Shorts\nB: Con thổi cho đã rồi ba đút mẹ ăn, dỗi hờn cả thế giới luôn! 😤🤣 #TikTokFamily\nC: Đừng đùa với công chúa nhỏ khi cô ấy đang đói! 🍗👶 #ReelsViral", font_name='Segoe UI', font_size=11)

    p_h4 = doc.add_heading(level=2)
    p_h4.add_run("4. Caption & Hashtag 3 Tầng")
    add_clean_paragraph(doc, "Tội nghiệp công chúa nhỏ của ba mẹ ghê chưa! Thổi cho nguội đã đời xong ba quay xe đút mẹ cắn một miếng to đùng! Nhìn cái mặt dỗi với pha liếc xéo mà ba mẹ cười muốn xỉu! Cả nhà có con nhỏ dỗi hờn dễ thương vậy không nè? 🍗👶🤣\n\n#Shorts #TikTokFamily #ReelsViral #HanDaily #LinhMitDaily #BabyFunny #EmBeHaiHuoc #GiaDinhHaiHuoc #CuLuaGaRan #TrendingShorts", font_name='Segoe UI', font_size=10.5)

    doc.save(os.path.join(output_dir, "04_METADATA_SEO.docx"))

def export_all():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    output_dir = os.path.join(base_dir, "export_batch")
    os.makedirs(output_dir, exist_ok=True)
    export_file_00(output_dir)
    export_file_01(output_dir)
    export_file_02(output_dir)
    export_file_03(output_dir)
    export_file_04(output_dir)
    print(f"[OK] Đã xuất thành công trọn bộ 5 file .docx chuẩn Tool Batch tại: {output_dir}")

if __name__ == '__main__':
    export_all()
